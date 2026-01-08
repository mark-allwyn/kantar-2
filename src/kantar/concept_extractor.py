"""
LLM-based concept extractor for Kantar PPTX presentations.

Uses GPT-4 to extract structured concept information from PowerPoint slides,
handling various presentation formats and layouts.
"""

import json
from pathlib import Path
from typing import List, Optional, Dict
from pptx import Presentation
import logging

from ..parsers.concept_parser import Concept
from ..llm.client import LLMClient

logger = logging.getLogger(__name__)


class ConceptExtractor:
    """
    LLM-based extractor for concept presentations.

    Uses GPT-4 to intelligently parse PPTX slides and extract structured
    concept information regardless of presentation format.
    """

    EXTRACTION_PROMPT = """You are analyzing a PowerPoint slide that describes a product concept for a survey study.

Extract the following information from the slide text:

1. **Concept Name**: The title or name of the concept (usually prominent at the top)
2. **Description**: A brief description of what the concept is about
3. **Price**: Any pricing information mentioned (e.g., "$5", "500 CZK", "Free")
4. **Features**: List of key features, benefits, or characteristics mentioned
5. **Occasion**: If the concept is tied to a specific occasion (e.g., Christmas, Birthday, Valentine's)
6. **Category**: Product category if mentioned (e.g., lottery, gaming, scratch card)

SLIDE TEXT:
{slide_text}

Respond with a JSON object in this exact format:
{{
    "name": "concept name",
    "description": "brief description",
    "price": "price or null",
    "features": ["feature 1", "feature 2"],
    "occasion": "occasion or null",
    "category": "category or null"
}}

If you cannot identify a value, use null. Be concise but accurate."""

    def __init__(self, model: str = "gpt-4o"):
        """
        Initialize the concept extractor.

        Args:
            model: LLM model to use for extraction
        """
        self.llm_client = LLMClient(model=model)
        self.cache_dir = Path("data/kantar-survey-source")

    def extract_from_pptx(self, pptx_path: Path, cache: bool = True) -> List[Concept]:
        """
        Extract concepts from a PowerPoint presentation.

        Args:
            pptx_path: Path to PPTX file
            cache: If True, cache results as JSON

        Returns:
            List of Concept objects

        Raises:
            FileNotFoundError: If PPTX file doesn't exist
            ValueError: If extraction fails
        """
        if not pptx_path.exists():
            raise FileNotFoundError(f"PPTX file not found: {pptx_path}")

        # Check if cached JSON exists
        cache_path = self._get_cache_path(pptx_path)
        if cache and cache_path.exists():
            logger.info(f"Loading concepts from cache: {cache_path}")
            return self._load_from_cache(cache_path)

        logger.info(f"Extracting concepts from: {pptx_path}")

        try:
            prs = Presentation(str(pptx_path))
        except Exception as e:
            raise ValueError(f"Failed to load PPTX file: {e}")

        concepts = []
        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1

            # Extract all text from slide
            slide_text = self._extract_slide_text(slide)

            if not slide_text.strip():
                logger.debug(f"Slide {slide_num}: No text found, skipping")
                continue

            # Use LLM to extract concept information
            try:
                concept = self._extract_concept_with_llm(
                    slide_num=slide_num,
                    slide_text=slide_text
                )

                if concept:
                    concepts.append(concept)
                    logger.info(f"Slide {slide_num}: Extracted concept '{concept.name}'")
                else:
                    logger.debug(f"Slide {slide_num}: Not a concept slide")

            except Exception as e:
                logger.warning(f"Slide {slide_num}: Extraction failed - {e}")
                continue

        logger.info(f"Extracted {len(concepts)} concepts from {pptx_path.name}")

        # Cache results
        if cache and concepts:
            self._save_to_cache(concepts, cache_path)

        return concepts

    def _extract_slide_text(self, slide) -> str:
        """Extract all text content from a slide."""
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    text_parts.append(text)

        return "\n".join(text_parts)

    def _extract_concept_with_llm(self, slide_num: int, slide_text: str) -> Optional[Concept]:
        """
        Use LLM to extract concept information from slide text.

        Args:
            slide_num: Slide number (for concept ID)
            slide_text: Raw text from slide

        Returns:
            Concept object or None if not a concept slide
        """
        # Quick heuristic: skip obviously non-concept slides
        if len(slide_text) < 20:
            return None

        # Skip title/intro slides
        skip_indicators = [
            'thank you', 'agenda', 'introduction', 'overview',
            'disclaimer', 'contact', 'appendix'
        ]
        if any(ind in slide_text.lower() for ind in skip_indicators):
            return None

        # Build prompt
        prompt = self.EXTRACTION_PROMPT.format(slide_text=slide_text)

        # Call LLM
        try:
            response = self.llm_client.generate_response(
                system_prompt="You are a helpful assistant that extracts structured information from text.",
                user_prompt=prompt,
                temperature=0.0  # Deterministic for consistency
            )

            # Parse JSON response
            response_text = response.strip()

            # Handle markdown code blocks if present
            if response_text.startswith("```"):
                lines = response_text.split("\n")
                response_text = "\n".join(lines[1:-1])  # Remove first and last line

            data = json.loads(response_text)

            # Validate that this looks like a concept
            if not data.get('name') or data['name'].lower() in ['none', 'null', 'n/a']:
                return None

            # Build description from components
            description_parts = []
            if data.get('description'):
                description_parts.append(data['description'])
            if data.get('price'):
                description_parts.append(f"Price: {data['price']}")
            if data.get('occasion'):
                description_parts.append(f"Occasion: {data['occasion']}")
            if data.get('features'):
                description_parts.append("Features:")
                description_parts.extend([f"- {f}" for f in data['features']])

            description = "\n".join(description_parts)

            return Concept(
                id=f"Concept {slide_num}",
                name=data['name'],
                description=description,
                price=data.get('price'),
                features=data.get('features', []),
                occasion=data.get('occasion'),
                full_text=slide_text,
                image_description=data.get('category')  # Repurpose for category
            )

        except json.JSONDecodeError as e:
            logger.error(f"Failed to parse LLM response as JSON: {e}")
            logger.error(f"Response: {response}")
            return None
        except Exception as e:
            logger.error(f"LLM extraction failed: {e}")
            return None

    def _get_cache_path(self, pptx_path: Path) -> Path:
        """Get path for cached concepts JSON file."""
        return pptx_path.parent / "concepts.json"

    def _save_to_cache(self, concepts: List[Concept], cache_path: Path):
        """Save concepts to cache file."""
        data = [
            {
                'id': c.id,
                'name': c.name,
                'description': c.description,
                'price': c.price,
                'features': c.features,
                'occasion': c.occasion,
                'full_text': c.full_text,
                'image_description': c.image_description
            }
            for c in concepts
        ]

        try:
            with open(cache_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            logger.info(f"Cached concepts to: {cache_path}")
        except Exception as e:
            logger.warning(f"Failed to cache concepts: {e}")

    def _load_from_cache(self, cache_path: Path) -> List[Concept]:
        """Load concepts from cache file."""
        try:
            with open(cache_path, 'r', encoding='utf-8') as f:
                data = json.load(f)

            concepts = [
                Concept(
                    id=c['id'],
                    name=c['name'],
                    description=c['description'],
                    price=c.get('price'),
                    features=c.get('features', []),
                    occasion=c.get('occasion'),
                    full_text=c.get('full_text', ''),
                    image_description=c.get('image_description')
                )
                for c in data
            ]

            return concepts
        except Exception as e:
            logger.warning(f"Failed to load cache: {e}")
            return []


def main():
    """CLI entry point for testing the extractor."""
    import sys
    from .study_catalog import StudyCatalog

    logging.basicConfig(level=logging.INFO)

    # Test with iGaming US market
    catalog = StudyCatalog()
    study = catalog.get_study('61405445-01')

    if not study:
        print("Study not found")
        sys.exit(1)

    market = study.get_market('US')
    if not market or not market.pptx_files:
        print("US market or PPTX file not found")
        sys.exit(1)

    extractor = ConceptExtractor()
    concepts = extractor.extract_from_pptx(market.pptx_files[0])

    print(f"\n=== Extracted Concepts ===")
    print(f"Total: {len(concepts)}\n")

    for concept in concepts:
        print(f"{concept.id}: {concept.name}")
        if concept.price:
            print(f"  Price: {concept.price}")
        if concept.occasion:
            print(f"  Occasion: {concept.occasion}")
        if concept.features:
            print(f"  Features: {len(concept.features)}")
        print()


if __name__ == "__main__":
    main()
